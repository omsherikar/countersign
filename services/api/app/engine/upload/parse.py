"""PARSE: read a CSV, XLSX or JSON rows into raw rows and check the header (API_CONTRACT §A2.2,
§A4.16).

Headers are matched after trimming, lower-casing and turning spaces and hyphens into
underscores, through the college's saved column mapping first (`mapping.apply`). Unknown
headers are reported in `ignored_columns`. A header that appears twice, or a missing
required header, rejects the file: UPLOAD_HEADER_UNMAPPED when unknown headers could be
mapped to the missing columns, else UPLOAD_HEADER_INVALID. Blank rows are skipped and not
counted; `row_no` is the spreadsheet row number (header row 1), or the 1-based position for
rows.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Iterator, Mapping
from datetime import date, datetime
from pathlib import Path
from typing import Any

from app.core.config import get_settings
from app.engine.upload import mapping as column_mapping
from app.engine.upload.types import (
    UPLOAD_COLUMNS,
    FileKind,
    ParsedUpload,
    RawRow,
    UploadFileError,
)

_KNOWN = frozenset(UPLOAD_COLUMNS)

UNREADABLE = "The file could not be read as CSV or XLSX."
NOT_UTF8 = "The file is not UTF-8 text. Save it as CSV UTF-8 and upload it again."
EMPTY = "The file has a header but no rows."


normalise_header = column_mapping.normalise_header


def file_kind(filename: str) -> FileKind | None:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return "csv"
    if lower.endswith(".xlsx"):
        return "xlsx"
    return None


def too_many_rows(limit: int) -> UploadFileError:
    return UploadFileError(
        "TOO_MANY_ROWS", f"The file has more than {limit:,} rows. Split it and upload each part."
    )


def cell_text(value: Any) -> str:
    """A cell or JSON value as text: dates as YYYY-MM-DD, whole numbers without a decimal point."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _check_header(
    cells: list[Any], saved: Mapping[str, str] | None
) -> tuple[list[tuple[int, str]], list[str], list[tuple[str, str]]]:
    """(column index, known column) pairs, the ignored header texts and the (header, column)
    pairs read through the saved mapping; raises on a bad header."""
    header_map = column_mapping.apply([cell_text(c) for c in cells], saved)
    missing, duplicates = header_map.missing, header_map.duplicates
    if missing and header_map.unknown and not duplicates:
        raise UploadFileError(
            "UPLOAD_HEADER_UNMAPPED",
            column_mapping.UNMAPPED,
            column_mapping.unmapped_details(header_map),
        )
    if missing or duplicates:
        if missing:
            message = f"The file is missing required columns: {', '.join(missing)}."
        else:
            message = f"The column {duplicates[0]} appears more than once."
        raise UploadFileError(
            "UPLOAD_HEADER_INVALID",
            message,
            {"missing_columns": missing, "duplicate_columns": duplicates},
        )
    return sorted(header_map.columns.items()), header_map.unknown, header_map.mapped


def _rows(
    records: Iterable[tuple[int, list[Any]]], columns: list[tuple[int, str]], max_rows: int
) -> list[RawRow]:
    rows: list[RawRow] = []
    for row_no, cells in records:
        texts = [cell_text(c) for c in cells]
        if not any(texts):
            continue
        if len(rows) >= max_rows:
            raise too_many_rows(max_rows)
        values = dict.fromkeys(UPLOAD_COLUMNS, "")
        for index, key in columns:
            if index < len(texts):
                values[key] = texts[index]
        rows.append(RawRow(row_no=row_no, values=values))
    if not rows:
        raise UploadFileError("UPLOAD_EMPTY", EMPTY)
    return rows


def _from_records(
    records: Iterator[tuple[int, list[Any]]],
    kind: FileKind,
    max_rows: int,
    saved: Mapping[str, str] | None,
) -> ParsedUpload:
    first = next(records, None)
    columns, ignored, mapped = _check_header(first[1] if first else [], saved)
    return ParsedUpload(
        rows=_rows(records, columns, max_rows),
        ignored_columns=ignored,
        kind=kind,
        column_mapping=[{"header": h, "column": c} for h, c in mapped],
    )


def _csv_records(data: bytes) -> Iterator[tuple[int, list[Any]]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise UploadFileError("UNSUPPORTED_FILE_TYPE", NOT_UTF8) from None
    if "\x00" in text:
        raise UploadFileError("UNSUPPORTED_FILE_TYPE", UNREADABLE)
    try:
        records = list(csv.reader(io.StringIO(text, newline="")))
    except csv.Error:
        raise UploadFileError("UNSUPPORTED_FILE_TYPE", UNREADABLE) from None
    return enumerate(records, start=1)


def _xlsx_records(data: bytes) -> Iterator[tuple[int, list[Any]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # not a zip, not a workbook, or a workbook openpyxl cannot open
        raise UploadFileError("UNSUPPORTED_FILE_TYPE", UNREADABLE) from None
    try:
        if not workbook.worksheets:
            raise UploadFileError("UNSUPPORTED_FILE_TYPE", UNREADABLE)
        sheet = workbook.worksheets[0]
        records = [
            (row_no, list(cells))
            for row_no, cells in enumerate(sheet.iter_rows(values_only=True), start=1)
        ]
    except UploadFileError:
        raise
    except Exception:
        raise UploadFileError("UNSUPPORTED_FILE_TYPE", UNREADABLE) from None
    finally:
        workbook.close()
    return iter(records)


def parse_upload(
    data: bytes | Path,
    filename: str,
    *,
    max_rows: int | None = None,
    mapping: Mapping[str, str] | None = None,
) -> ParsedUpload:
    """Read a CSV (UTF-8, BOM allowed) or XLSX (first sheet), reading the header through the
    college's saved column mapping. Raises UploadFileError."""
    limit = max_rows if max_rows is not None else get_settings().upload_max_rows
    if isinstance(data, Path):
        data = data.read_bytes()
    kind = file_kind(filename)
    if kind == "csv":
        return _from_records(_csv_records(data), "csv", limit, mapping)
    if kind == "xlsx":
        return _from_records(_xlsx_records(data), "xlsx", limit, mapping)
    raise UploadFileError("UNSUPPORTED_FILE_TYPE", UNREADABLE)


def parse_rows(
    rows: list[dict[str, Any]],
    *,
    max_rows: int | None = None,
    mapping: Mapping[str, str] | None = None,
) -> ParsedUpload:
    """JSON rows from POST /uploads/rows: the template columns as keys, `null` means blank;
    keys are read through the college's saved column mapping as file headers are."""
    limit = max_rows if max_rows is not None else get_settings().upload_max_rows
    saved = {normalise_header(k): v for k, v in (mapping or {}).items()}
    ignored: dict[str, None] = {}
    used: dict[str, str] = {}
    parsed: list[RawRow] = []
    for position, row in enumerate(rows, start=1):
        values = dict.fromkeys(UPLOAD_COLUMNS, "")
        seen: set[str] = set()
        for raw_key, raw_value in row.items():
            key = normalise_header(str(raw_key))
            column = saved.get(key, key)
            if column in seen:
                raise UploadFileError(
                    "UPLOAD_HEADER_INVALID",
                    f"The column {column} appears more than once.",
                    {"missing_columns": [], "duplicate_columns": [column]},
                )
            seen.add(column)
            if column in _KNOWN:
                values[column] = cell_text(raw_value)
                if key in saved:
                    used.setdefault(str(raw_key).strip(), column)
            elif str(raw_key).strip():
                ignored.setdefault(str(raw_key).strip(), None)
        if not any(values.values()):
            continue
        if len(parsed) >= limit:
            raise too_many_rows(limit)
        parsed.append(RawRow(row_no=position, values=values))
    if not parsed:
        raise UploadFileError("UPLOAD_EMPTY", EMPTY)
    return ParsedUpload(
        rows=parsed,
        ignored_columns=list(ignored),
        kind="json",
        column_mapping=[{"header": h, "column": c} for h, c in used.items()],
    )
