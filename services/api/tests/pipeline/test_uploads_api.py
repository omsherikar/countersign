"""Upload endpoints: multipart and JSON uploads, scoping, errors and templates (§A2.4 to §A2.9)."""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.config import get_settings
from app.db.session import session_factory
from app.engine import pipeline
from app.engine.upload.types import UPLOAD_COLUMNS
from app.models.audit import AuditLog
from app.models.upload import UploadBatch
from tests.conftest import login_as
from tests.pipeline.conftest import ROW, csv_bytes, institution, student_row
from tests.pipeline.fakes import FakeEngines

FORM = {"academic_year": "2026-27", "declaration": "true"}


@pytest.fixture
async def api(world, spool, monkeypatch) -> AsyncIterator[None]:
    """Uploads run inline with the fake engines."""
    monkeypatch.setattr(get_settings(), "upload_inline", True)
    monkeypatch.setattr(pipeline, "default_engines", FakeEngines)
    yield


@pytest.fixture
async def clients(api) -> AsyncIterator[dict[str, httpx.AsyncClient]]:
    emails = {
        "admin": "admin@aicte.demo",
        "abc": "registrar@abc.edu.demo",
        "xyz": "registrar@xyz.edu.demo",
        "west": "ro.west@aicte.demo",
        "south": "ro.south@aicte.demo",
        "university": "coe@univ.demo",
    }
    out = {key: await login_as(email) for key, email in emails.items()}
    yield out
    for client in out.values():
        await client.aclose()


async def post_file(
    client: httpx.AsyncClient,
    data: bytes,
    filename: str = "1-ABC-0001_2026-27.csv",
    form: dict | None = None,
) -> httpx.Response:
    content_type = "text/csv" if filename.endswith(".csv") else "application/octet-stream"
    return await client.post(
        "/uploads",
        data=FORM if form is None else form,
        files={"file": (filename, data, content_type)},
    )


def error(response: httpx.Response) -> tuple[int, str, str]:
    body = response.json()["error"]
    return response.status_code, body["code"], body["message"]


async def test_multipart_upload_runs_and_reports(clients, spool):
    abc = clients["abc"]
    response = await post_file(abc, csv_bytes([ROW, student_row(1, mobile="12")]))
    assert response.status_code == 202, response.text
    batch = response.json()["batch"]
    assert response.headers["location"] == f"/uploads/{batch['id']}"
    assert (batch["status"], batch["stage"], batch["progress_pct"]) == ("QUEUED", None, 0)
    assert (batch["source"], batch["kind"]) == ("FILE", "INITIAL")
    assert batch["filename"] == "1-ABC-0001_2026-27.csv"
    assert batch["institution_name"] == "ABC College of Engineering"
    assert batch["uploaded_by"]["name"] == "Registrar, ABC College of Engineering"
    assert list(batch["counts"]) == [
        "rows", "accepted", "errors", "warnings", "created", "updated", "unchanged", "withdrawn",
        "transferred", "seated", "released", "unseated", "flags_raised", "flags_updated",
        "versions", "records",
    ]  # fmt: skip
    assert batch["sources"] == [] and batch["same_file_as"] is None

    done = (await abc.get(f"/uploads/{batch['id']}")).json()
    assert (done["status"], done["stage"], done["progress_pct"]) == ("DONE", "DONE", 100)
    assert done["counts"]["rows"] == 2 and done["counts"]["created"] == 2
    # the done card says how the batch's records stand (API_CONTRACT §A4.16)
    assert sum(done["counts"]["records"].values()) == 2
    assert done["counts"]["warnings"] == 1 and done["total_rows"] == 2
    assert set(done["timings_ms"]) == {s.value for s in pipeline.STAGES}
    assert done["uploaded_at"].endswith("Z") and done["finished_at"].endswith("Z")
    assert list(spool.iterdir()) == []

    async with session_factory()() as session:
        stored = await session.get(UploadBatch, uuid.UUID(batch["id"]))
        entry = (
            await session.execute(
                select(AuditLog).where(
                    AuditLog.entity_type == "upload_batch", AuditLog.action == "CREATE"
                )
            )
        ).scalar_one()
    assert stored.declared_at is not None and stored.declaration_text.startswith("I declare")
    assert len(stored.file_sha256) == 64
    assert entry.actor_id is not None and entry.after["source"] == "FILE"
    assert entry.after["filename"] == "1-ABC-0001_2026-27.csv"

    again = await post_file(abc, csv_bytes([ROW, student_row(1, mobile="12")]))
    second = again.json()["batch"]
    assert second["kind"] == "UPDATE"
    assert second["same_file_as"] == {"batch_id": batch["id"], "uploaded_at": batch["uploaded_at"]}
    finished = (await abc.get(f"/uploads/{second['id']}")).json()
    assert finished["counts"]["unchanged"] == 2 and finished["counts"]["created"] == 0


async def test_xlsx_upload_and_admin_upload_for_a_college(clients):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(UPLOAD_COLUMNS))
    sheet.append([ROW.get(c, "") for c in UPLOAD_COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    xyz = await institution("1-XYZ-0002")
    data = buffer.getvalue()
    form = dict(FORM, institution_id=str(xyz.id))
    response = await post_file(clients["admin"], data, "roll.XLSX", form)
    assert response.status_code == 202, response.text
    batch = response.json()["batch"]
    assert batch["institution_id"] == str(xyz.id) and batch["filename"] == "roll.XLSX"
    done = (await clients["admin"].get(f"/uploads/{batch['id']}")).json()
    # CSE is not a programme of XYZ
    assert done["status"] == "DONE" and done["counts"]["errors"] == 1


async def test_upload_validation_errors(clients, monkeypatch):
    abc, admin = clients["abc"], clients["admin"]
    data = csv_bytes([ROW])

    response = await post_file(abc, data, form={"academic_year": "2026-27"})
    assert error(response) == (422, "VALIDATION_ERROR", "Tick the declaration to upload.")
    assert response.json()["error"]["details"]["fields"] == [
        {"field": "declaration", "message": "Tick the declaration to upload."}
    ]
    response = await post_file(abc, data, form={"academic_year": "2027-28", "declaration": "true"})
    assert error(response) == (
        422,
        "VALIDATION_ERROR",
        "Your college has no approvals for 2027-28.",
    )
    response = await abc.post("/uploads", data=FORM, files={"other": ("x.csv", data)})
    assert error(response)[2] == "Choose a file to upload."
    assert error(await post_file(admin, data)) == (
        422,
        "VALIDATION_ERROR",
        "Choose the college this upload is for.",
    )
    unknown = dict(FORM, institution_id="00000000-0000-4000-8000-000000000000")
    assert error(await post_file(admin, data, form=unknown)) == (
        404,
        "NOT_FOUND",
        "Institution not found.",
    )
    xyz = await institution("1-XYZ-0002")
    other = dict(FORM, institution_id=str(xyz.id))
    assert error(await post_file(abc, data, form=other))[0] == 404
    own = dict(FORM, institution_id=str((await institution()).id))
    assert (await post_file(abc, data, form=own)).status_code == 202

    assert error(await post_file(abc, data, "roll.pdf")) == (
        415,
        "UNSUPPORTED_FILE_TYPE",
        "Only .csv and .xlsx files can be uploaded.",
    )
    assert error(await post_file(abc, "name,dob\nJosé,x\n".encode("latin-1")))[0:2] == (
        415,
        "UNSUPPORTED_FILE_TYPE",
    )
    header = await post_file(abc, b"apaar_id,name,name\n1,2,3\n")
    assert error(header) == (
        422,
        "UPLOAD_HEADER_INVALID",
        "The file is missing required columns: dob, programme_code, year_of_study, "
        "admission_route, category.",
    )
    assert header.json()["error"]["details"] == {
        "missing_columns": [
            "dob",
            "programme_code",
            "year_of_study",
            "admission_route",
            "category",
        ],
        "duplicate_columns": ["name"],
    }
    assert error(await post_file(abc, csv_bytes([])))[:2] == (422, "UPLOAD_EMPTY")

    monkeypatch.setattr(get_settings(), "upload_max_rows", 1)
    assert error(await post_file(abc, csv_bytes([ROW, student_row(1)]))) == (
        413,
        "TOO_MANY_ROWS",
        "The file has more than 1 rows. Split it and upload each part.",
    )
    monkeypatch.setattr(get_settings(), "upload_max_bytes", 2 * 1024 * 1024)
    assert error(await post_file(abc, b"x" * (2 * 1024 * 1024 + 1))) == (
        413,
        "FILE_TOO_LARGE",
        "The file is larger than 2 MB. Split it by programme and upload each part.",
    )

    as_json = await abc.post("/uploads", json={"academic_year": "2026-27"})
    assert as_json.status_code == 415
    assert as_json.json()["error"]["message"] == "Request bodies must be multipart/form-data."


async def test_upload_roles(clients):
    for key in ("west", "university"):
        response = await post_file(clients[key], csv_bytes([ROW]))
        assert response.status_code == 403
        rows = await clients[key].post("/uploads/rows", json={"rows": []})
        assert rows.status_code == 403


async def test_upload_in_progress_and_queue(clients, monkeypatch):
    monkeypatch.setattr(get_settings(), "upload_inline", False)
    sent = []
    monkeypatch.setattr("app.worker.enqueue_batch", lambda batch_id: sent.append(batch_id))
    first = await post_file(clients["abc"], csv_bytes([ROW]))
    assert first.status_code == 202 and first.json()["batch"]["status"] == "QUEUED"
    assert [str(b) for b in sent] == [first.json()["batch"]["id"]]

    busy = await post_file(clients["abc"], csv_bytes([student_row(2)]))
    assert error(busy) == (
        409,
        "UPLOAD_IN_PROGRESS",
        "An upload for this college and year is still being processed.",
    )
    assert busy.json()["error"]["details"] == {"batch_id": first.json()["batch"]["id"]}
    # another college is not blocked
    xyz_row = dict(ROW, programme_code="ECE", apaar_id="")
    assert (await post_file(clients["xyz"], csv_bytes([xyz_row]))).status_code == 202


async def test_a_batch_that_cannot_be_queued_fails_at_once(clients, monkeypatch):
    monkeypatch.setattr(get_settings(), "upload_inline", False)

    def broken(batch_id):
        raise ConnectionError("redis is down")

    monkeypatch.setattr("app.worker.enqueue_batch", broken)
    response = await post_file(clients["abc"], csv_bytes([ROW]))
    assert response.status_code == 202
    batch = (await clients["abc"].get(f"/uploads/{response.json()['batch']['id']}")).json()
    assert batch["status"] == "FAILED" and batch["stage"] == "PARSE"
    assert batch["error"].startswith("Processing stopped while reading the file.")
    assert (await post_file(clients["abc"], csv_bytes([ROW]))).status_code == 202


async def test_json_rows(clients):
    body = {
        "academic_year": "2026-27",
        "institution_id": None,
        "declaration": True,
        "label": "SIS sync 28 Oct",
        "rows": [
            dict(ROW, year_of_study=1, sis_id="A1"),
            {"name": None},
            dict(student_row(3), status="WITHDRAWN", status_date="2026-10-20"),
        ],
    }
    response = await clients["abc"].post("/uploads/rows", json=body)
    assert response.status_code == 202, response.text
    batch = response.json()["batch"]
    assert (batch["source"], batch["filename"], batch["status"]) == (
        "API",
        "SIS sync 28 Oct",
        "QUEUED",
    )
    done = (await clients["abc"].get(f"/uploads/{batch['id']}")).json()
    assert done["status"] == "DONE" and done["ignored_columns"] == ["sis_id"]
    assert done["counts"]["rows"] == 2 and done["counts"]["withdrawn"] == 1
    errors = (await clients["abc"].get(f"/uploads/{batch['id']}/errors")).json()
    assert [(e["row_no"], e["code"]) for e in errors["items"]] == [(3, "NOT_PREVIOUSLY_ENROLLED")]

    unnamed = dict(body, label=None, rows=[dict(student_row(4))])
    later = (await clients["abc"].post("/uploads/rows", json=unnamed)).json()["batch"]
    assert later["filename"] == "rows.json"

    assert (
        error(await clients["abc"].post("/uploads/rows", json=dict(body, declaration=False)))[2]
        == "Tick the declaration to upload."
    )
    assert (
        error(await clients["abc"].post("/uploads/rows", json=dict(body, label="x" * 121)))[2]
        == "Label must be at most 120 characters."
    )
    assert error(await clients["abc"].post("/uploads/rows", json=dict(body, rows=[])))[:2] == (
        422,
        "UPLOAD_EMPTY",
    )
    too_many = dict(body, rows=[student_row(n) for n in range(1, 5002)])
    assert error(await clients["abc"].post("/uploads/rows", json=too_many)) == (
        413,
        "TOO_MANY_ROWS",
        "The file has more than 5,000 rows. Split it and upload each part.",
    )


async def test_list_and_detail_are_scoped(clients):
    abc_batch = (await post_file(clients["abc"], csv_bytes([ROW]))).json()["batch"]
    xyz_row = dict(ROW, programme_code="ECE")
    xyz_batch = (await post_file(clients["xyz"], csv_bytes([xyz_row]))).json()["batch"]

    def ids(response: httpx.Response) -> list[str]:
        return [b["id"] for b in response.json()["items"]]

    everything = await clients["admin"].get("/uploads")
    assert ids(everything) == [xyz_batch["id"], abc_batch["id"]]  # newest first
    assert everything.json()["total"] == 2 and everything.json()["page_size"] == 25
    abc = await institution()
    only_abc = await clients["admin"].get("/uploads", params={"institution_id": str(abc.id)})
    assert ids(only_abc) == [abc_batch["id"]]
    assert ids(await clients["west"].get("/uploads")) == [abc_batch["id"]]
    assert ids(await clients["south"].get("/uploads")) == [xyz_batch["id"]]
    # an institution's own list ignores institution_id
    mine = await clients["xyz"].get("/uploads", params={"institution_id": str(abc.id)})
    assert ids(mine) == [xyz_batch["id"]]
    assert ids(await clients["admin"].get("/uploads", params={"status": "FAILED"})) == []
    assert ids(await clients["admin"].get("/uploads", params={"academic_year": "2025-26"})) == []
    assert (await clients["university"].get("/uploads")).status_code == 403

    for key in ("xyz", "south"):
        response = await clients[key].get(f"/uploads/{abc_batch['id']}")
        assert error(response) == (404, "NOT_FOUND", "Upload not found.")
        assert (await clients[key].get(f"/uploads/{abc_batch['id']}/errors")).status_code == 404
        assert (await clients[key].get(f"/uploads/{abc_batch['id']}/errors.csv")).status_code == 404
    assert (await clients["west"].get(f"/uploads/{abc_batch['id']}")).status_code == 200
    assert (await clients["university"].get(f"/uploads/{abc_batch['id']}")).status_code == 403


async def test_row_errors_and_csv(clients):
    rows = [
        ROW,
        student_row(1, dob="21/03/2007", mobile="12"),
        student_row(2, apaar_id="12345678441"),
        student_row(3, apaar_id=ROW["apaar_id"]),
    ]
    batch = (await post_file(clients["abc"], csv_bytes(rows), "Roll 2026 (final).csv")).json()[
        "batch"
    ]
    url = f"/uploads/{batch['id']}/errors"
    page = (await clients["abc"].get(url)).json()
    assert page["total"] == 4 and page["page"] == 1
    assert [(e["row_no"], e["field"], e["code"], e["level"]) for e in page["items"]] == [
        (3, "dob", "DOB_FORMAT", "ERROR"),
        (3, "mobile", "MOBILE_FORMAT", "WARNING"),
        (4, "apaar_id", "APAAR_FORMAT", "WARNING"),
        (5, "apaar_id", "DUPLICATE_IN_FILE", "ERROR"),
    ]
    assert page["items"][2] == {
        "row_no": 4,
        "level": "WARNING",
        "field": "apaar_id",
        "code": "APAAR_FORMAT",
        "message": "APAAR must be 12 digits",
        "value_masked": "<11 characters>",
    }
    warnings = (await clients["abc"].get(url, params={"level": "WARNING"})).json()
    assert warnings["total"] == 2
    dupes = (await clients["abc"].get(url, params={"code": "duplicate_in_file"})).json()
    assert [e["message"] for e in dupes["items"]] == ["Same APAAR appears on row 2"]
    paged = (await clients["abc"].get(url, params={"page": 2, "page_size": 3})).json()
    assert [e["row_no"] for e in paged["items"]] == [5]

    done = (await clients["abc"].get(f"/uploads/{batch['id']}")).json()["counts"]
    assert (done["accepted"], done["errors"], done["warnings"]) == (2, 2, 1)

    download = await clients["west"].get(f"{url}.csv")
    assert download.status_code == 200
    assert download.headers["content-type"] == "text/csv; charset=utf-8"
    assert download.headers["content-disposition"] == (
        'attachment; filename="Roll 2026 _final__errors.csv"'
    )
    text = download.content.decode("utf-8")
    assert text.startswith("﻿row_no,level,field,code,message,value_masked\r\n")
    assert "5,ERROR,apaar_id,DUPLICATE_IN_FILE,Same APAAR appears on row 2,XXXX XXXX 4417" in text
    assert len(text.strip().splitlines()) == 5


async def test_templates(clients):
    response = await clients["west"].get("/uploads/template.csv")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        'attachment; filename="countersign_upload_template.csv"'
    )
    assert response.text.strip() == ",".join(UPLOAD_COLUMNS)
    xlsx = await clients["abc"].get("/uploads/template.xlsx")
    assert xlsx.headers["content-disposition"].endswith('countersign_upload_template.xlsx"')
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == ["students"]
    assert [c.value for c in next(workbook.active.iter_rows())] == list(UPLOAD_COLUMNS)
    assert (await clients["university"].get("/uploads/template.csv")).status_code == 403
