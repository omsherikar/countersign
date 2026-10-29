"""The renderers as pure functions: every file is labelled synthetic and keeps the columns."""

from __future__ import annotations

import base64
import csv
import io
import re
import zlib
from datetime import UTC, datetime

from openpyxl import load_workbook

from app.reports import render
from app.reports.registry import REPORTS, ReportScope

AT = datetime(2026, 10, 29, 5, 50, tzinfo=UTC)
WEST = ReportScope("REGION", None, None, "Western Region", "WEST")


def pdf_pages(data: bytes) -> int:
    return len(re.findall(rb"/Type /Page[^s]", data))


def pdf_text(data: bytes) -> str:
    """The page content streams, decoded: enough to find the words a page prints."""
    out = []
    for stream in re.findall(rb"stream\r?\n(.*?)\r?\n?endstream", data, re.S):
        raw = stream.strip()
        if raw.endswith(b"~>"):
            raw = base64.a85decode(raw[:-2])
        try:
            out.append(zlib.decompress(raw).decode("latin-1"))
        except zlib.error:
            continue
    return "".join(out)


def coverage_rows(n: int) -> list[dict]:
    return [
        {
            "academic_year": "2026-27",
            "aicte_permanent_id": f"1-T-{i:04d}",
            "institution_name": f"College {i:02d}",
            "state_code": "MH",
            "institution_type": "GOVERNMENT",
            "records": 10,
            "countersigned": 4,
            "awaiting_evidence": 5,
            "under_review": 1,
            "disputed": 0,
            "countersigned_pct": 40.0,
            "independent_joining_record": 4,
            "independent_credits": 0,
            "independent_student": 2,
        }
        for i in range(n)
    ]


def test_csv_starts_with_the_synthetic_line_then_the_wire_columns():
    spec = REPORTS["completeness"]
    rows = [
        {
            "academic_year": "2026-27",
            "aicte_permanent_id": "1-ABC-0001",
            "institution_name": "ABC College of Engineering",
            "state_code": "MH",
            "joined": 140,
            "uploaded_and_joined": 132,
            "missing": 8,
            "uploaded_not_joined": 3,
            "completeness_pct": 94.28571,
            "deadline": "2026-09-30",
            "first_upload_on": "2026-09-22",
            "deadline_state": "ON_TIME",
            "case_status": None,
        }
    ]
    data = render.to_csv(spec, rows)
    assert data.startswith("﻿".encode())
    lines = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert lines[0] == [render.SYNTHETIC]
    assert lines[1] == list(spec.columns)
    assert lines[2][spec.columns.index("completeness_pct")] == "94.3"
    assert lines[2][-1] == "" and lines[2][spec.columns.index("deadline")] == "2026-09-30"


def test_xlsx_has_the_report_sheet_and_an_about_sheet():
    spec = REPORTS["countersign-coverage"]
    book = load_workbook(
        io.BytesIO(render.to_xlsx(spec, coverage_rows(3), WEST, "2026-27", generated_at=AT))
    )
    assert book.sheetnames == ["Countersign coverage", "About"]
    sheet = book["Countersign coverage"]
    assert sheet["A1"].value == render.SYNTHETIC
    assert [c.value for c in sheet[2]] == list(spec.columns)
    assert sheet.max_row == 5
    about = {row[0]: row[1] for row in book["About"].iter_rows(values_only=True)}
    assert about == {
        "Title": "Countersign coverage",
        "Scope": "Western Region",
        "Academic year": "2026-27",
        "Generated": "29 Oct 2026, 11:20 IST",
        "Data": "Synthetic demo data. Every number comes from the Countersign demo dataset.",
    }


def test_pdf_is_one_page_with_25_rows_a_total_and_the_rest_counted():
    spec = REPORTS["countersign-coverage"]
    data = render.to_pdf(
        spec, coverage_rows(40), WEST, "2026-27", generated_at=AT, by="an AICTE officer"
    )
    assert data.startswith(b"%PDF") and pdf_pages(data) == 1
    text = pdf_text(data)
    assert "Synthetic data. Every college, student and number here is generated" in text
    assert "and 15 more rows in the CSV or XLSX" in text
    assert "College 24" in text and "College 25" not in text
    assert "Generated 29 Oct 2026, 11:20 IST by an AICTE officer. Page 1 of 1." in text
    totals = render._totals(spec.columns, coverage_rows(40))
    assert (totals["records"], totals["countersigned"], totals["countersigned_pct"]) == (
        400,
        160,
        40.0,
    )
    assert "institution_name" not in totals


def test_pdf_with_no_rows_still_renders():
    spec = REPORTS["by-state"]
    data = render.to_pdf(spec, [], WEST, "2026-27", generated_at=AT)
    assert pdf_pages(data) == 1 and "No data for 2026-27 in this scope." in pdf_text(data)


def test_summary_pdf_fits_one_page():
    data = {
        "kpis": {
            "institutions": 50,
            "approved_seats": 18960,
            "admitted": 4120,
            "records": 10342,
            "countersigned_pct": 61.2,
            "cases_open": 214,
            "joined": 4210,
            "missing_from_upload": 120,
            "completeness_pct": None,
            "colleges_late": 4,
            "colleges_overdue": 2,
            "signals_new": 6,
        },
        "records_by_status": [
            {"status": "COUNTERSIGNED", "count": 6329, "pct": 61.2},
            {"status": "AWAITING_EVIDENCE", "count": 3110, "pct": 30.1},
            {"status": "UNDER_REVIEW", "count": 895, "pct": 8.7},
            {"status": "DISPUTED", "count": 8, "pct": 0.1},
        ],
        "cases_by_type": [
            {"case_type": "EXCESS_ADMISSION", "case_name": "Excess admission", "open": 9}
        ]
        * 15,
        "top_institutions": [
            {
                "institution_name": f"College {i}",
                "aicte_permanent_id": f"1-T-{i:04d}",
                "cases_open": 40 - i,
                "under_review": 30 - i,
                "disputed": 1,
            }
            for i in range(10)
        ],
        "late_and_overdue": [
            {
                "institution_name": f"Late {i}",
                "aicte_permanent_id": f"1-L-{i}",
                "deadline_state": "LATE",
            }
            for i in range(14)
        ],
        "deadline": "2026-09-30",
    }
    pdf = render.summary_pdf(
        data, WEST, "2026-27", generated_at=AT, by="the Western regional office"
    )
    assert pdf_pages(pdf) == 1
    text = pdf_text(pdf)
    assert "Countersign summary, Western Region, 2026-27" in text
    assert "and 4 more in the completeness report" in text
    assert "Not computed" in text
