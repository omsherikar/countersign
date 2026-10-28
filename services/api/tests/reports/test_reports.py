"""Reports over the mini seed: every report in every format, scoped for AICTE and a regional
officer, and the same counts the dashboard and the institutions list give."""

from __future__ import annotations

import csv
import io
import uuid
from collections.abc import AsyncIterator
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, func, select

from app.db.session import session_factory
from app.models.audit import AuditLog
from app.models.credential import InstitutionScore
from app.models.enums import AuditAction, EnrolmentStatus, RecordStatus
from app.models.student import Enrolment
from app.reports import render
from app.reports.registry import REPORTS
from tests.conftest import institution_id
from tests.reports.test_render import pdf_pages, pdf_text

FORMATS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def read_csv(data: bytes) -> list[dict]:
    reader = csv.reader(io.StringIO(data.decode("utf-8-sig")))
    assert next(reader) == [render.SYNTHETIC]
    header = next(reader)
    return [dict(zip(header, row, strict=True)) for row in reader]


@pytest.fixture
async def abc_reconciled(seeded) -> AsyncIterator[str]:
    """ABC College's stored reconciliation: 25 joined, 20 uploaded, 5 missing."""
    abc = await institution_id("1-ABC-0001")
    async with session_factory()() as session:
        session.add(
            InstitutionScore(
                institution_id=uuid.UUID(abc),
                academic_year="2026-27",
                completeness_pct=Decimal("80.00"),
                components={
                    "joined": 25,
                    "uploaded_and_joined": 20,
                    "missing": 5,
                    "uploaded_not_joined": 1,
                    "missing_students": [
                        {
                            "name": "Sneha Kulkarni",
                            "roll_no": "MHCET-2026-007731",
                            "programme_code": "CSE",
                            "joined_on": "2026-08-04",
                            "round": 2,
                            "route": "COUNSELLING",
                            "apaar_last4": "8812",
                        }
                    ],
                    "published": True,
                    "source_unavailable": False,
                    "first_upload_at": None,
                },
            )
        )
        await session.commit()
    yield abc
    async with session_factory()() as session:
        await session.execute(delete(InstitutionScore))
        await session.commit()


async def test_only_officers_list_and_download_reports(admin, registrar_abc, university):
    listed = (await admin.get("/reports")).json()["items"]
    assert [r["name"] for r in listed] == [
        "by-state",
        "by-institution-type",
        "cases-by-type-and-outcome",
        "countersign-coverage",
        "completeness",
        "summary",
    ]
    assert listed[-1] == {
        "name": "summary",
        "title": "One-page summary",
        "formats": ["pdf"],
        "columns": [],
    }
    assert listed[0]["columns"][:3] == ["academic_year", "state_code", "state_name"]
    for client in (registrar_abc, university):
        assert (await client.get("/reports")).status_code == 403
        assert (await client.get("/reports/by-state.csv")).status_code == 403


async def test_every_report_in_every_format_for_aicte_and_a_region(admin, ro_west):
    for client, slug in ((admin, "national"), (ro_west, "WEST")):
        for spec in REPORTS.values():
            for fmt in spec.formats:
                response = await client.get(f"/reports/{spec.name}.{fmt}")
                assert response.status_code == 200, (spec.name, fmt, response.text)
                assert response.headers["content-type"] == FORMATS[fmt]
                assert response.headers["content-disposition"] == (
                    f'attachment; filename="countersign_{spec.name}_2026-27_{slug}.{fmt}"'
                )
                if fmt == "csv":
                    rows = read_csv(response.content)
                    assert all(list(r) == list(spec.columns) for r in rows)
                elif fmt == "xlsx":
                    book = load_workbook(io.BytesIO(response.content))
                    assert book.sheetnames == [spec.title, "About"]
                else:
                    assert pdf_pages(response.content) == 1
                    assert "Synthetic data." in pdf_text(response.content)

    national = read_csv((await admin.get("/reports/by-state.csv")).content)
    west = read_csv((await ro_west.get("/reports/by-state.csv")).content)
    assert [r["state_code"] for r in national] == ["KA", "MH", "TN"]
    assert [r["state_code"] for r in west] == ["MH"]
    west_coverage = read_csv((await ro_west.get("/reports/countersign-coverage.csv")).content)
    assert [r["aicte_permanent_id"] for r in west_coverage] == ["1-ABC-0001"]
    # a regional officer cannot widen its scope
    ignored = await ro_west.get("/reports/by-state.csv", params={"region": "SOUTH"})
    assert [r["state_code"] for r in read_csv(ignored.content)] == ["MH"]
    outside = await ro_west.get("/reports/by-state.csv", params={"state": "KA"})
    assert outside.status_code == 422
    one_state = await admin.get("/reports/countersign-coverage.xlsx", params={"state": "tn"})
    assert one_state.headers["content-disposition"].endswith('_2026-27_TN.xlsx"')

    async with session_factory()() as session:
        entry = (
            await session.execute(
                select(AuditLog)
                .where(AuditLog.entity_type == "report", AuditLog.action == AuditAction.READ)
                .order_by(AuditLog.seq.desc())
                .limit(1)
            )
        ).scalar_one()
    assert entry.after == {"name": "countersign-coverage", "format": "xlsx", "scope": "TN"}


async def test_unknown_reports_and_formats_are_not_found(admin):
    for path in ("/reports/nope.csv", "/reports/summary.csv", "/reports/by-state.docx"):
        response = await admin.get(path)
        assert response.status_code == 404
        assert response.json()["error"] == {
            "code": "REPORT_NOT_FOUND",
            "message": "Report not found.",
            "details": {},
        }
    assert (await admin.get("/reports/by-state.csv", params={"state": "ZZ"})).status_code == 422


async def test_the_dashboard_reports_and_institutions_count_the_same_records(admin, ro_west):
    async with session_factory()() as session:
        stored = dict(
            (
                await session.execute(
                    select(Enrolment.record_status, func.count())
                    .where(
                        Enrolment.academic_year == "2026-27",
                        Enrolment.status == EnrolmentStatus.ACTIVE,
                    )
                    .group_by(Enrolment.record_status)
                )
            ).all()
        )
    expected = {s.value: stored.get(s, 0) for s in RecordStatus}
    kpis = (await admin.get("/dashboard/summary")).json()["kpis"]
    assert kpis["records"] == expected

    by_state = read_csv((await admin.get("/reports/by-state.csv")).content)
    key = {"AWAITING_EVIDENCE": "awaiting_evidence", "COUNTERSIGNED": "countersigned"}
    key |= {"UNDER_REVIEW": "under_review", "DISPUTED": "disputed"}
    for status, column in key.items():
        assert sum(int(r[column]) for r in by_state) == expected[status]
    assert sum(int(r["records"]) for r in by_state) == kpis["enrolments"]
    assert sum(int(r["cases_open"]) for r in by_state) == kpis["flags_open"]

    coverage = read_csv((await admin.get("/reports/countersign-coverage.csv")).content)
    listed = {
        i["aicte_permanent_id"]: i for i in (await admin.get("/institutions")).json()["items"]
    }
    for row in coverage:
        institution = listed[row["aicte_permanent_id"]]
        assert int(row["countersigned"]) == institution["records"]["COUNTERSIGNED"]
        assert int(row["records"]) == institution["enrolments"]
        assert (float(row["countersigned_pct"]) if row["countersigned_pct"] else None) == (
            institution["countersigned_pct"]
        )

    west = (await ro_west.get("/dashboard/summary")).json()["kpis"]
    west_rows = read_csv((await ro_west.get("/reports/by-state.csv")).content)
    assert sum(int(r["records"]) for r in west_rows) == west["enrolments"]

    cases = read_csv((await admin.get("/reports/cases-by-type-and-outcome.csv")).content)
    assert len(cases) == 15
    assert (
        sum(int(r["open"]) + int(r["docs_requested"]) + int(r["escalated"]) for r in cases)
        == (kpis["flags_open"])
    )
    assert next(r for r in cases if r["case_type"] == "LATE_REFILL")["case_name"] == "Late refill"


async def test_completeness_comes_from_the_stored_reconciliation(
    abc_reconciled, admin, registrar_abc
):
    rows = read_csv((await admin.get("/reports/completeness.csv")).content)
    assert rows[0]["aicte_permanent_id"] == "1-ABC-0001"  # lowest completeness first
    abc = rows[0]
    assert (abc["joined"], abc["uploaded_and_joined"], abc["missing"]) == ("25", "20", "5")
    assert (abc["completeness_pct"], abc["deadline"], abc["deadline_state"]) == (
        "80.0",
        "2026-09-30",
        "LATE",
    )
    assert all(r["completeness_pct"] == "" for r in rows[1:])  # nulls last

    kpis = (await admin.get("/dashboard/summary")).json()["kpis"]
    assert (kpis["joined"], kpis["missing_from_upload"], kpis["completeness_pct"]) == (25, 5, 80.0)
    by_state = read_csv((await admin.get("/reports/by-state.csv")).content)
    mh = next(r for r in by_state if r["state_code"] == "MH")
    assert (mh["joined"], mh["missing_from_upload"], mh["completeness_pct"]) == ("25", "5", "80.0")

    college = (await registrar_abc.get("/dashboard/summary")).json()["completeness"]
    assert college["completeness_pct"] == 80.0 and college["published"] is True
    assert college["missing_students"][0]["apaar_masked"] == "XXXX XXXX 8812"
    detail = (await admin.get(f"/institutions/{abc_reconciled}")).json()
    assert (
        detail["completeness"]["missing"] == 5 and "missing_students" not in detail["completeness"]
    )
    assert (detail["completeness_pct"], detail["missing_from_upload"]) == (80.0, 5)

    ordered = (await admin.get("/institutions", params={"sort": "completeness_asc"})).json()
    assert ordered["items"][0]["aicte_permanent_id"] == "1-ABC-0001"
    late = (await admin.get("/institutions", params={"deadline_state": "LATE"})).json()
    on_time = (await admin.get("/institutions", params={"deadline_state": "ON_TIME"})).json()
    assert (late["total"], on_time["total"]) == (3, 0)
    by_share = (await admin.get("/institutions", params={"sort": "countersigned_asc"})).json()
    shares = [i["countersigned_pct"] for i in by_share["items"]]
    assert shares == sorted(shares, key=lambda v: (v is None, v))
