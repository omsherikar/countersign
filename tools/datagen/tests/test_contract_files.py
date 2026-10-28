"""The section E file contract (with Addendum 2): names, columns, shapes, demo users."""

import datetime as dt
import re

from openpyxl import load_workbook

from datagen.config import load_datagen_config
from datagen.students import UPLOAD_COLUMNS

ELIGIBILITY_COLUMNS = ["tfw_income_certified", "nri_sponsored", "ciwg"]
CONTRACT_COLUMNS = [
    "apaar_id", "name", "dob", "gender", "mobile", "email", "address", "nationality",
    "passport_no", "programme_code", "academic_year", "year_of_study", "admission_route",
    "counselling_ref", "category", "consent_status", "status", "status_date", "admission_date",
]  # fmt: skip

DEMO_USERS = [
    ("admin@aicte.demo", "AICTE_ADMIN"),
    ("ro.west@aicte.demo", "RO_OFFICER"),
    ("ro.south@aicte.demo", "RO_OFFICER"),
    ("registrar@abc.edu.demo", "INSTITUTION"),
    ("registrar@xyz.edu.demo", "INSTITUTION"),
    ("student@demo", "STUDENT"),
    ("coe@univ.demo", "UNIVERSITY"),
]


def test_every_contract_file_exists(small):
    for name in [
        "institutions.json", "users.json", "ground_truth.json", "manifest.json", "summary.json",
        "mockgov/apaar_registry.jsonl", "mockgov/counselling.jsonl",
        "mockgov/abc_credits.jsonl", "mockgov/sii.jsonl", "mockgov/joining_records.jsonl",
        "history.json",
    ]:  # fmt: skip
        assert (small.path / name).is_file(), name
    insts = small.json("institutions.json")["institutions"]
    for inst in insts:
        aid = inst["aicte_permanent_id"]
        assert (small.path / "uploads" / f"{aid}_2026-27.csv").is_file()
        assert (small.path / "uploads" / f"{aid}_2026-27.xlsx").is_file()
    updates = small.update_files()
    assert updates
    for name in updates:
        assert re.fullmatch(r"uploads/[^_]+_2026-27_update1\.csv", name)
        assert (small.path / name.replace(".csv", ".xlsx")).is_file()
    assert len(list((small.path / "uploads").glob("*.csv"))) == len(insts) + len(updates)
    assert len(list((small.path / "uploads").glob("*.xlsx"))) == len(insts) + len(updates)


def test_csv_columns_match_contract(small):
    """The contract's nineteen columns in order, then the optional eligibility markers."""
    assert UPLOAD_COLUMNS == CONTRACT_COLUMNS + ELIGIBILITY_COLUMNS
    for aid in ["1-ABC-0001", "1-XYZ-0002"]:
        assert small.upload_columns(aid) == UPLOAD_COLUMNS


def test_xlsx_mirrors_csv(small):
    rows = small.upload("1-ABC-0001")
    wb = load_workbook(small.path / "uploads" / "1-ABC-0001_2026-27.xlsx", read_only=True)
    ws = wb["students"]
    values = list(ws.iter_rows(values_only=True))
    assert list(values[0]) == UPLOAD_COLUMNS
    assert len(values) - 1 == len(rows)
    first = values[1]
    assert first[0] == "123456784417"
    assert first[1] == "Ravi Kumar Sharma"
    assert str(first[11]) == "1"
    assert first[18] == rows[0]["admission_date"]
    for name in small.update_files():
        csv_rows = small.file(name)
        wb = load_workbook(small.path / name.replace(".csv", ".xlsx"), read_only=True)
        values = list(wb["students"].iter_rows(values_only=True))
        assert list(values[0]) == UPLOAD_COLUMNS
        assert len(values) - 1 == len(csv_rows)
        for cells, row in zip(values[1:], csv_rows, strict=True):
            # A read-only sheet leaves out trailing empty cells, as the api's reader expects.
            cells = list(cells) + [None] * (len(UPLOAD_COLUMNS) - len(cells))
            assert [("" if c is None else str(c)) for c in cells] == [
                row[c] for c in UPLOAD_COLUMNS
            ]


def admission_window(year_of_study: int) -> tuple[str, str]:
    start = 2026 - (year_of_study - 1)
    end = "2026-09-15" if start == 2026 else f"{start}-09-30"
    return f"{start}-07-01", end


def test_initial_files_leave_status_blank_and_date_every_admission(small):
    for aid, rows in small.all_uploads().items():
        for row in rows:
            assert row["status"] == "" and row["status_date"] == "", aid
            assert dt.date.fromisoformat(row["admission_date"])
            lo, hi = admission_window(int(row["year_of_study"]))
            assert lo <= row["admission_date"] <= hi, (aid, row["year_of_study"])


def test_institutions_json_shape(small):
    data = small.json("institutions.json")
    assert set(data) == {"regions", "states", "universities", "institutions"}
    assert [r["code"] for r in data["regions"]][3] == "WEST"
    assert {"code", "name", "hq_city"} == set(data["regions"][0])
    assert {"code", "name", "region"} == set(data["states"][0])
    assert {"code", "name", "state_code"} == set(data["universities"][0])
    inst = data["institutions"][0]
    assert set(inst) == {
        "aicte_permanent_id", "name", "type", "state_code", "university_code",
        "address_city", "established_year", "autonomous", "programmes",
    }  # fmt: skip
    prog = inst["programmes"][0]
    assert {"code", "name", "level", "discipline", "shift", "approvals"} <= set(prog)
    assert {"academic_year", "approved_intake", "versions"} == set(prog["approvals"][0])
    assert {"version", "approved_intake", "effective_from", "source"} == set(
        prog["approvals"][0]["versions"][0]
    )
    assert inst["aicte_permanent_id"] == "1-ABC-0001"
    assert inst["name"] == "ABC College of Engineering"
    assert inst["state_code"] == "MH"
    assert inst["university_code"] == "MSTU"
    xyz = data["institutions"][1]
    assert xyz["aicte_permanent_id"] == "1-XYZ-0002"
    assert xyz["name"] == "XYZ Institute of Technology"
    assert xyz["state_code"] == "TN"
    assert any(u["code"] == "MSTU" for u in data["universities"])
    assert len(data["institutions"]) == 50


def test_users_are_exactly_the_demo_set_plus_registrars(small):
    users = small.json("users.json")
    assert [(u["email"], u["role"]) for u in users[:7]] == DEMO_USERS
    assert all(u["password"] == "Countersign!2026" for u in users)
    by_email = {u["email"]: u for u in users}
    assert by_email["ro.west@aicte.demo"]["region"] == "WEST"
    assert by_email["ro.south@aicte.demo"]["region"] == "SOUTH"
    assert by_email["registrar@abc.edu.demo"]["institution_aicte_id"] == "1-ABC-0001"
    assert by_email["registrar@xyz.edu.demo"]["institution_aicte_id"] == "1-XYZ-0002"
    assert by_email["student@demo"]["student_ref"] == "123456784417"
    assert by_email["student@demo"]["name"] == "Ravi Kumar Sharma"
    assert by_email["coe@univ.demo"]["university_code"] == "MSTU"
    registrars = users[7:]
    assert len(registrars) == 48
    assert all(u["role"] == "INSTITUTION" for u in registrars)
    assert len({u["email"] for u in users}) == len(users)
    inst_ids = {i["aicte_permanent_id"] for i in small.json("institutions.json")["institutions"]}
    assert {u["institution_aicte_id"] for u in users if u["role"] == "INSTITUTION"} == inst_ids


def test_mockgov_jsonl_shapes(small):
    reg = small.jsonl("apaar_registry.jsonl")
    assert {"apaar_id", "name", "dob", "gender"} == set(reg[0])
    assert all(re.fullmatch(r"\d{12}", r["apaar_id"]) for r in reg)
    assert len({r["apaar_id"] for r in reg}) == len(reg)
    demo = next(r for r in reg if r["apaar_id"] == "123456784417")
    assert demo["name"] == "Ravi Kumar Sharma" and demo["dob"] == "2007-03-21"

    cou = small.jsonl("counselling.jsonl")
    assert set(cou[0]) == {
        "state", "year", "name", "dob", "roll_no", "programme", "institution_aicte_id",
        "round", "category", "allotment_ref",
    }  # fmt: skip
    assert len({c["roll_no"] for c in cou}) == len(cou)
    assert any(c["allotment_ref"] == "MHCET-2026-R1-004512" for c in cou)

    abc = small.jsonl("abc_credits.jsonl")
    assert set(abc[0]) == {
        "apaar_id", "academic_year", "sem", "credits", "uploaded_at", "result", "issuer",
    }  # fmt: skip
    assert {"type", "code", "name"} == set(abc[0]["issuer"])
    assert all(a["uploaded_at"].endswith("Z") for a in abc)
    assert all(a["result"] in {"PASS", "FAIL", "PENDING"} for a in abc)
    assert len({(a["apaar_id"], a["academic_year"], a["sem"]) for a in abc}) == len(abc)

    sii = small.jsonl("sii.jsonl")
    assert set(sii[0]) == {
        "passport_no", "name", "nationality", "institution_aicte_id", "programme", "academic_year",
    }  # fmt: skip
    assert all(re.fullmatch(r"[A-Z]{2}", s["nationality"]) for s in sii)
    assert len({s["passport_no"] for s in sii}) == len(sii)


def test_manifest_and_summary(small):
    manifest = small.json("manifest.json")
    assert set(manifest) == {
        "size", "seed", "generated_at", "counts", "demo_unanswered_notices", "deck_sample",
    }  # fmt: skip
    assert manifest["size"] == "small" and manifest["seed"] == 42
    counts = manifest["counts"]
    assert counts["institutions"] == 50
    assert counts["students"] == sum(len(r) for r in small.all_uploads().values())
    assert counts["apaar_registry"] == len(small.jsonl("apaar_registry.jsonl"))
    assert counts["upload_files"] == len(small.all_uploads())
    assert counts["update_files"] == len(small.update_files())
    assert counts["update_rows"] == sum(len(small.file(f)) for f in small.update_files())
    summary = small.json("summary.json")
    assert summary["counts"] == counts
    assert summary["xlsx_written"] is True
    assert summary["timings_seconds"]["total"] > 0


def test_student_count_near_target(small):
    target = load_datagen_config()["sizes"]["small"]["target_students"]
    students = small.json("manifest.json")["counts"]["students"]
    assert 0.85 * target <= students <= 1.2 * target
